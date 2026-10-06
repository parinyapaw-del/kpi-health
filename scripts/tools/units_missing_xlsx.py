#!/usr/bin/env python3
# local one-off (not part of the pipeline; the Actions bot never runs it) - needs openpyxl only
"""Write docs/units_missing_<year>.xlsx from data/lookup/units.json -> `missing` (reporting units the MOPH GIS
registry does not know, located by the fallback rule) — offline, no network.

  python3 scripts/tools/units_missing_xlsx.py [--year 2569] [--out docs/units_missing_2569.xlsx]

Sheet "units": one row per unit (target descending) — hospcode, province, district code+name and fallback tambon
code+name (`maxTargetArea`), override tambon (units.json `overrides`, the HDC-confirmed location that wins over the
fallback), DSPM target, target per tambon, GIS error. Sheet "districts": units per fallback district.
`missing[h].target/areas` are summed over the years in sites/<site>.json when units.json was written; --year only
labels the columns.
"""
import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]      # repo root (scripts/tools/ -> ../..)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--year", default="2569", help="fiscal year label of the target columns")
    ap.add_argument("--out", default=None, help="default docs/units_missing_<year>.xlsx")
    a = ap.parse_args()
    out = Path(a.out) if a.out else ROOT / "docs" / f"units_missing_{a.year}.xlsx"
    reg = json.loads((ROOT / "data/lookup/units.json").read_text(encoding="utf-8"))
    lk = json.loads((ROOT / "data/lookup/areas.json").read_text(encoding="utf-8"))
    prov = {k: v["name"] for k, v in lk["provinces"].items()}
    dist, sub = lk["districts"], lk["subdistricts"]
    overrides = reg.get("overrides", {})
    missing = sorted(reg["missing"].items(), key=lambda kv: (-kv[1].get("target", 0), kv[0]))

    wb = Workbook()
    ws = wb.active
    ws.title = "units"
    head = ["hospcode", "รหัสจังหวัด", "จังหวัด", "รหัสอำเภอ", "อำเภอ (fallback)", "รหัสตำบล fallback",
            "ตำบล fallback (เป้ามากสุด)", "override ตำบล (HDC)", f"เป้า DSPM {a.year} รวม",
            "จำนวนตำบลที่หน่วยส่งข้อมูล", "รายละเอียดเป้าต่อตำบล", "gisError", "gisFailedAt"]
    ws.append(head)
    by_dist: dict = {}
    for h, m in missing:
        t6 = m.get("maxTargetArea") or ""
        d4, p2 = t6[:4], t6[:2]
        ov = overrides.get(h, {}).get("tambon", "")
        areas = m.get("areas", {})
        detail = "; ".join(f"{sub.get(c, c)}={n}" for c, n in sorted(areas.items(), key=lambda kv: (-kv[1], kv[0])))
        ws.append([h, p2, prov.get(p2, ""), d4, dist.get(d4, ""), t6, sub.get(t6, ""),
                   f"{ov} {sub.get(ov, '')}".strip() if ov else "", m.get("target", 0), len(areas), detail,
                   m.get("gisError", ""), m.get("gisFailedAt", "")])
        b = by_dist.setdefault(d4, {"n": 0, "target": 0, "codes": []})
        b["n"] += 1
        b["target"] += m.get("target", 0)
        b["codes"].append(h)

    ws2 = wb.create_sheet("districts")
    ws2.append(["ลำดับ", "จังหวัด", "อำเภอ", "รหัสอำเภอ", "จำนวนหน่วยนอกทะเบียน", f"เป้า DSPM {a.year} ของหน่วยเหล่านี้",
                "รหัสหน่วย"])
    for i, (d4, b) in enumerate(sorted(by_dist.items(), key=lambda kv: (-kv[1]["target"], kv[0])), 1):
        ws2.append([i, prov.get(d4[:2], ""), dist.get(d4, ""), d4, b["n"], b["target"], ", ".join(sorted(b["codes"]))])

    for w in (ws, ws2):
        for c in w[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        w.freeze_panes = "A2"
        for col in w.columns:
            width = max(len(str(c.value or "")) for c in col)
            w.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, width * 0.9), 70)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out.relative_to(ROOT) if out.is_relative_to(ROOT) else out}: {len(missing)} units, "
          f"{len(by_dist)} districts, {sum(1 for _, m in missing if 'gisError' in m)} with gisError")


if __name__ == "__main__":
    main()

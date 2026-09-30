"""population_report — (1) reconcile the API caches with the user's HDC exports in "manual data by user/",
(2) tabulate the 0-5 denominators of the home province, (3) write an Excel workbook with live formulas.

Outputs
  docs/population_tables_angthong.md                      generated tables (reconciliation + denominators)
  data/prepared/angthong/population_denominators.xlsx     one sheet per (year, level) + reconciliation sheet
The narrative analysis lives in docs/POPULATION_DENOMINATORS_ANGTHONG.md (hand-written, not overwritten here).
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import build_site as B
from .indicators import denom05, pop, pyramid, typearea
from .indicators.popcommon import group_sums, scope_keyfns

ROOT = Path(__file__).resolve().parents[1]
MANUAL = ROOT / "manual data by user"
PREPARED = ROOT / "data" / "prepared"
DOCS = ROOT / "docs"

LEVEL_BY_HEADER = {"เขตสุขภาพ": ("country", "TH"), "จังหวัด": ("region", None), "อำเภอ": ("province", None),
                   "ตำบล": ("district", None), "หมู่": ("village", None)}
FOLDER = {"sexage": "ประชากรจำแนกเพศ กลุ่มอายุรายปี", "pyramid": "ปิรามิดประชากรจำแนกเพศ กลุ่มอายุ",
          "typearea": "ประชากรแยกตามหน่วยบริการและชนิดการอยู่อาศัย TYPEAREA"}


# ------------------------------------------------------------------ manual Excel readers
def _find_folder(name: str) -> Path | None:
    for p in MANUAL.iterdir():
        if p.is_dir() and p.name.strip() == name:
            return p
    return None


def _xlsx_files(folder: Path):
    return sorted(f for f in folder.iterdir() if f.suffix.lower() in (".xlsx", ".xlxs") and not f.name.startswith("~$"))


def _open(path: Path):
    try:
        return openpyxl.load_workbook(path, data_only=True).active
    except Exception as e:      # noqa: BLE001  (.xlxs typo etc.)
        return e


def read_sexage(path):
    ws = _open(path)
    if not hasattr(ws, "iter_rows"):
        return {"error": str(ws)}
    header = str(ws.cell(1, 1).value or "").strip()
    rows = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] is None:
            continue
        vals = list(row[1:70])
        mf = {}
        for b in range(22):
            mf[f"male_g{b + 1}"] = vals[3 * b] or 0
            mf[f"female_g{b + 1}"] = vals[3 * b + 1] or 0
        rows[str(row[0]).strip()] = mf
    return {"header": header, "rows": rows}


def read_typearea(path):
    ws = _open(path)
    if not hasattr(ws, "iter_rows"):
        return {"error": str(ws)}
    header = str(ws.cell(1, 1).value or "").strip()
    rows = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[0] is None:
            continue
        rows[str(row[0]).strip()] = {f: (row[i + 1] or 0) for i, f in enumerate(typearea.RAW_FIELDS)}
    return {"header": header, "rows": rows}


def read_pyramid(path):
    ws = _open(path)
    if not hasattr(ws, "iter_rows"):
        return {"error": str(ws)}
    rows = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        rows[str(row[0]).strip()] = {"male": row[1] or 0, "female": row[2] or 0, "total": row[3] or 0}
    return {"header": "ช่วงอายุ", "rows": rows}


# ------------------------------------------------------------------ reconciliation
def _name_to_code(ctx, level, scope):
    lk = ctx["lookup"]
    if level == "country":
        return {v: k for k, v in lk["regions"].items()}
    if level == "region":
        return {v["name"]: k for k, v in lk["provinces"].items() if str(v["region"]) == scope}
    if level == "province":
        return {n: c for c, n in lk["districts"].items() if c[:2] == scope}
    return {n: c for c, n in lk["subdistricts"].items() if c[:4] == scope}


def _guess_scope(ctx, level, names):
    """Excel files carry no codes: find the home-path scope whose child names match the row labels."""
    home = ctx["home"]
    cands = {"country": ["TH"], "region": [str(home["region"])], "province": [home["province"]],
             "district": [home["district"]]}[level]
    for sc in cands:
        m = _name_to_code(ctx, level, sc)
        if names <= set(m):
            return sc, m
    return None, {}


def reconcile(ctx, years, log=print) -> list[dict]:
    """Compare every manual Excel with every cached year; return one record per (file, year)."""
    recs = []
    fx = _find_folder(FOLDER["sexage"])
    ft = _find_folder(FOLDER["typearea"])
    fp = _find_folder(FOLDER["pyramid"])
    for kind, folder in (("sexage", fx), ("typearea", ft), ("pyramid", fp)):
        if folder is None:
            continue
        for f in _xlsx_files(folder):
            rel = str(f.relative_to(MANUAL))
            xl = {"sexage": read_sexage, "typearea": read_typearea, "pyramid": read_pyramid}[kind](f)
            if "error" in xl:
                recs.append({"file": rel, "kind": kind, "status": "unreadable", "detail": xl["error"]})
                continue
            names = {n for n in xl["rows"] if n != "รวม"}
            if kind == "pyramid":
                recs += _reconcile_pyramid(ctx, years, rel, xl, names)
                continue
            level = LEVEL_BY_HEADER.get(xl["header"], (None, None))[0]
            if level in (None, "village"):
                recs.append({"file": rel, "kind": kind, "status": "skipped", "level": level or xl["header"],
                             "detail": "ระดับหมู่บ้าน/ไม่รู้จัก header — ไม่เทียบ (API มีถึงระดับหมู่ แต่ตารางนี้ไม่ได้ทำ view)"})
                continue
            scope, m = _guess_scope(ctx, level, names)
            if scope is None:
                recs.append({"file": rel, "kind": kind, "status": "skipped", "level": level,
                             "detail": f"ชื่อแถวไม่ตรง lookup: {sorted(names)[:3]}..."})
                continue
            plugin = pop if kind == "sexage" else typearea
            for y in years:
                cache = plugin.load_cache(y)
                if cache is None:
                    continue
                sums = group_sums(cache, scope_keyfns(ctx, level, scope))
                fields = plugin.RAW_FIELDS
                n_rows = n_cells = n_diff = 0
                max_rel = 0.0
                examples = []
                for name in sorted(names):
                    code = m[name]
                    a = sums.get(code, dict.fromkeys(fields, 0))
                    x = xl["rows"][name]
                    n_rows += 1
                    row_diff = False
                    for fld in fields:
                        n_cells += 1
                        if a[fld] != x[fld]:
                            n_diff += 1
                            row_diff = True
                            base = max(x[fld], 1)
                            max_rel = max(max_rel, abs(a[fld] - x[fld]) / base * 100)
                    if row_diff and len(examples) < 3:
                        ta, tx = sum(a.values()), sum(x.values())
                        examples.append(f"{name}: API {ta:,} vs Excel {tx:,}")
                recs.append({"file": rel, "kind": kind, "level": level, "scope": scope, "year": y,
                             "status": "match" if n_diff == 0 else "diff", "rows": n_rows, "cells": n_cells,
                             "cells_diff": n_diff, "max_rel_pct": round(max_rel, 2), "examples": examples,
                             "asOf": cache["asOf"]})
    return recs


def _reconcile_pyramid(ctx, years, rel, xl, names):
    """Pyramid Excel has no area header: match its totals against the home province / district / a subdistrict."""
    recs = []
    home = ctx["home"]
    prov = home["province"]
    groups = [g for g in xl["rows"] if g != "รวม"]
    for y in years:
        cache = pyramid.load_cache(y, prov)
        if cache is None:
            continue
        best = None
        cands = [("province", prov, lambda a6, u6: "x"),
                 ("district", home["district"], lambda a6, u6: "x" if a6[:4] == home["district"] else None)]
        for s in B.child_codes(ctx, "district", home["district"]):
            cands.append(("subdistrict", s, lambda a6, u6, s=s: "x" if u6 == s else None))
        for level, scope, kf in cands:
            n_diff = 0
            for g in groups:
                api = pyramid.group_totals(cache, kf, g).get("x", 0)
                if api != xl["rows"][g]["total"]:
                    n_diff += 1
            if best is None or n_diff < best[0]:
                best = (n_diff, level, scope)
        n_diff, level, scope = best
        tot_api = sum(pyramid.group_totals(cache, [c for c in cands if c[1] == scope][0][2], g).get("x", 0)
                      for g in groups)
        tot_xl = sum(xl["rows"][g]["total"] for g in groups)
        recs.append({"file": rel, "kind": "pyramid", "level": level, "scope": scope, "year": y,
                     "status": "match" if n_diff == 0 else "diff", "rows": len(groups), "cells": len(groups),
                     "cells_diff": n_diff, "max_rel_pct": round(abs(tot_api - tot_xl) / max(tot_xl, 1) * 100, 2),
                     "examples": [f"รวมทุกอายุ: API {tot_api:,} vs Excel {tot_xl:,}"], "asOf": cache["asOf"],
                     "detail": "จับคู่พื้นที่จากผลรวมที่ใกล้ที่สุด (ไฟล์ปิรามิดไม่มีชื่อพื้นที่)"})
    return recs


# ------------------------------------------------------------------ tables / excel
def _load_prepared(site_id, ind, year, level, scope):
    p = PREPARED / site_id / B.dataset_file((ind, year, level, scope))
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


def _md_table(headers, rows):
    out = ["| " + " | ".join(headers) + " |", "|" + "|".join("---:" if i else "---" for i in range(len(headers))) + "|"]
    for r in rows:
        out.append("| " + " | ".join(r) + " |")
    return "\n".join(out)


def _fmt(v, t):
    if v is None:
        return "–"
    return f"{v:,.1f}" if t == "pct" else f"{v:,}"


DENOM_COLS = [("pop_lt1", "<1 ปี"), ("pop_1_4", "1-4 ปี"), ("pop_0_4", "0-4 ปี HDC"), ("pyramid_0_4", "ปิรามิด 0-4"),
              ("pop_0_5_est", "0-5 ประมาณ"), ("dspm_target", "เป้า DSPM"), ("dspm_screened", "คัดกรอง DSPM"),
              ("cov_population", "Coverage (1)"), ("cov_expected", "Coverage (3)"), ("midyear_all", "ปชก.ทุกอายุ"),
              ("pct_dspm_target", "เป้า DSPM/0-4 %"), ("pct_cov_pop", "Cov(1)/0-4 %"),
              ("pct_cov_pop_05", "Cov(1)/0-5 %"), ("pct_child_share", "0-4/ทุกอายุ %")]


def denominator_tables(site, ctx, years):
    types = {k: t for k, _, t in denom05.TABLE_KEYS}
    home = ctx["home"]
    scopes = [("region", str(home["region"])), ("province", home["province"]), ("district", home["district"]),
              ("country", "TH")]
    md = []
    sheets = []
    for y in years:
        for level, scope in scopes:
            d = _load_prepared(site["site"], "denom05", y, level, scope)
            if d is None:
                continue
            rows = d["rows"] + [dict(d["total"], name=f"รวม {d['scope']['name']}")]
            md.append(f"\n### ตัวหาร 0–5 ปี · ปีงบ {y} · {d['scope']['name']} (ราย{B.child_level(level) and {'region': 'เขต', 'province': 'จังหวัด', 'district': 'อำเภอ', 'subdistrict': 'ตำบล'}[B.child_level(level)]}) · ข้อมูล ณ {d['asOf']}\n")
            md.append(_md_table(["พื้นที่"] + [lab for _, lab in DENOM_COLS],
                                [[r["name"]] + [_fmt(r["metrics"][k], types[k]) for k, _ in DENOM_COLS] for r in rows]))
            sheets.append((y, level, scope, d, rows))
    return "\n".join(md), sheets


def write_xlsx(site, ctx, years, sheets, recs, path: Path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDEBF7")
    for y, level, scope, d, rows in sheets:
        ws = wb.create_sheet(f"{y}_{level}_{scope}"[:31])
        ws.append([f"ตัวหารเด็ก 0–5 ปี · ปีงบ {y} · {d['scope']['name']} · ข้อมูล ณ {d['asOf']} · แหล่ง: {d['source']}"])
        ws["A1"].font = bold
        count_cols = [c for c in DENOM_COLS if not c[0].startswith("pct_")]
        hdr = ["รหัส", "พื้นที่"] + [lab for _, lab in count_cols] + ["เป้า DSPM ÷ 0-4 (%)", "Coverage(1) ÷ 0-4 (%)",
                                                                    "Coverage(1) ÷ 0-5 ประมาณ (%)", "0-4 ÷ ทุกอายุ (%)",
                                                                    "ปิรามิด ÷ 0-4 (%)"]
        ws.append(hdr)
        for c in range(1, len(hdr) + 1):
            ws.cell(2, c).font = bold
            ws.cell(2, c).fill = fill
            ws.cell(2, c).alignment = Alignment(wrap_text=True, vertical="top")
        col = {k: get_column_letter(3 + i) for i, (k, _) in enumerate(count_cols)}
        for r in rows:
            m = r["metrics"]
            ws.append([r["code"], r["name"]] + [m[k] for k, _ in count_cols])
            i = ws.max_row
            f = lambda n, dd: f'=IF(OR({col[dd]}{i}="",{col[dd]}{i}=0,{col[n]}{i}=""),"",{col[n]}{i}/{col[dd]}{i}*100)'
            ws.cell(i, len(count_cols) + 3).value = f("dspm_target", "pop_0_4")
            ws.cell(i, len(count_cols) + 4).value = f("cov_population", "pop_0_4")
            ws.cell(i, len(count_cols) + 5).value = f("cov_population", "pop_0_5_est")
            ws.cell(i, len(count_cols) + 6).value = f("pop_0_4", "midyear_all")
            ws.cell(i, len(count_cols) + 7).value = f("pyramid_0_4", "pop_0_4")
            for c in range(len(count_cols) + 3, len(count_cols) + 8):
                ws.cell(i, c).number_format = "0.0"
            for c in range(3, len(count_cols) + 3):
                ws.cell(i, c).number_format = "#,##0"
        ws.cell(ws.max_row, 2).font = bold
        ws.column_dimensions["B"].width = 22
        for c in range(3, len(hdr) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 14
        ws.freeze_panes = "C3"
    # reconciliation sheet
    ws = wb.create_sheet("reconcile_manual_excel")
    hdr = ["ไฟล์ (manual data by user)", "ชนิด", "ระดับ", "scope", "ปี API", "ผล", "แถว", "เซลล์", "เซลล์ต่าง",
           "ต่างสูงสุด %", "API ณ", "ตัวอย่าง/หมายเหตุ"]
    ws.append(hdr)
    for c in range(1, len(hdr) + 1):
        ws.cell(1, c).font = bold
        ws.cell(1, c).fill = fill
    for r in recs:
        ws.append([r.get("file"), r.get("kind"), r.get("level"), r.get("scope"), r.get("year"), r.get("status"),
                   r.get("rows"), r.get("cells"), r.get("cells_diff"), r.get("max_rel_pct"), r.get("asOf"),
                   "; ".join(r.get("examples", [])) or r.get("detail", "")])
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["L"].width = 70
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def reconciliation_md(recs):
    rows = []
    for r in recs:
        if r["status"] in ("unreadable", "skipped"):
            rows.append([r["file"], r["kind"], r.get("level", "") or "", "", r["status"], "", r.get("detail", "")])
            continue
        rows.append([r["file"], r["kind"], f"{r['level']}/{r['scope']}", str(r["year"]),
                     "✅ ตรง 100%" if r["status"] == "match" else f"≠ {r['cells_diff']}/{r['cells']} เซลล์ (สูงสุด {r['max_rel_pct']}%)",
                     r["asOf"], "; ".join(r["examples"]) or r.get("detail", "")])
    return _md_table(["ไฟล์", "ชนิด", "ระดับ/scope", "ปี API", "ผล", "API ณ", "ตัวอย่าง"], rows)


def best_year_only(recs):
    """The Excel files carry no year: keep, per file, the cached year with the fewest differing cells."""
    best: dict = {}
    for r in recs:
        if r["status"] in ("unreadable", "skipped"):
            best.setdefault(r["file"], r)
            continue
        cur = best.get(r["file"])
        if cur is None or cur["status"] in ("unreadable", "skipped") or r["cells_diff"] < cur["cells_diff"]:
            best[r["file"]] = r
    return [best[f] for f in sorted(best)]


def run(site, ctx, years, log=print) -> list[str]:
    recs = best_year_only(reconcile(ctx, years, log))
    tables_md, sheets = denominator_tables(site, ctx, years)
    xlsx = PREPARED / site["site"] / "population_denominators.xlsx"
    write_xlsx(site, ctx, years, sheets, recs, xlsx)
    md = ["# ตารางประชากร HDC และตัวหารเด็ก 0–5 ปี (generated)",
          "",
          f"สร้างโดย `python3 scripts/population.py report` · อ่านคำอธิบาย/ข้อสรุปที่ `docs/POPULATION_DENOMINATORS_ANGTHONG.md`",
          f"· Excel สูตรสด: `{xlsx.relative_to(ROOT)}`",
          "",
          "## 1. เทียบ API (cache) กับ Excel ที่ export เองจาก HDC (`manual data by user/`)",
          "",
          "ไฟล์ Excel ไม่มีปี/รหัสพื้นที่ จึงเทียบกับ cache ทุกปีแล้วแสดงเฉพาะปีที่ต่างน้อยที่สุด (= ปีที่ export, 2569) ·",
          "ที่ต่างเล็กน้อยระดับเขต/ประเทศ = HDC ประมวลผลรายจังหวัดคนละเวลา (Excel export 30 ก.ย. 08:36–08:46 · API ดึง 09:2x)",
          "",
          reconciliation_md(recs),
          "",
          "## 2. ตัวหารเด็ก 0–5 ปี",
          "",
          "คอลัมน์: <1/1-4/0-4 = ประชากรกลางปี HDC (typearea 1+3) · ปิรามิด 0-4 = ตาราง s_person_pyramid (เฉพาะจังหวัดหลัก) ·",
          "0-5 ประมาณ = 0-4 + (5-9)/5 [INFERRED] · เป้า DSPM = 5 กลุ่มอายุรวม · Coverage (1) = เด็กปฐมวัย 0-5 ปี ตามตาราง coverage ·",
          "Coverage (3) = คาดประมาณเด็กล่าช้า (ความชุก 21.7%) · ปชก.ทุกอายุ = typearea 1c+3c",
          tables_md]
    out = DOCS / "population_tables_angthong.md"
    out.write_text("\n".join(md), encoding="utf-8")
    n_match = sum(1 for r in recs if r["status"] == "match")
    n_diff = sum(1 for r in recs if r["status"] == "diff")
    log(f"[report] {out.relative_to(ROOT)} + {xlsx.relative_to(ROOT)}  (reconcile: {n_match} match, {n_diff} diff, "
        f"{len(recs) - n_match - n_diff} skipped/unreadable)")
    for r in recs:
        if r["status"] == "match":
            log(f"  ✅ {r['file']}  ==  {r['level']}/{r['scope']} {r['year']}")
    return []
